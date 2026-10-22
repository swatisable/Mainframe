#!/usr/bin/env python3
"""
generate_mapping_workbook.py
============================

Builds ``CASPCFNJ_1x1_Field_Mapping.xlsx`` - a complete field-to-field mapping
for the COBOL program ``CASPCFNJ`` (PCF claim extract based on CAS2000 case
data).

The data model below is derived **directly** from the source program
``CASPCFNJ`` in this repository. Every mapping row carries the source-line
number of the MOVE / logic that populates it, so each entry can be traced back
to a real line of code.

IMPORTANT - copybooks are NOT present in this repository
--------------------------------------------------------
The program ``COPY``s many copybooks (FDPCF601/602, CLMPREFX, TPLPREFX,
NCTCASE, FDNJTPC5/4, FDNJ5INS, FDNJ4INS, ...). None of them are committed to
this repo. Therefore, for the main claim-extract output (SRCPCF-OUT), exact
PIC clauses and byte offsets cannot be read from source and are marked
``[FDPCF601 not in repo]`` (with an *inferred* PIC where the program logic makes
the type/length determinable). Picture/byte values that ARE fully determinable
from the program text (the 80-byte CASE-PCF-MATCH report defined inline in
WORKING-STORAGE) are given exactly.

Run:  python3 generate_mapping_workbook.py
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT_XLSX = "CASPCFNJ_1x1_Field_Mapping.xlsx"

# ---------------------------------------------------------------------------
# Styling helpers
# ---------------------------------------------------------------------------
HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(bold=True, color="FFFFFF", size=10)
SUB_FILL = PatternFill("solid", fgColor="D9E1F2")
XW_FILL = PatternFill("solid", fgColor="FCE4D6")   # crosswalk / override rows
DRV_FILL = PatternFill("solid", fgColor="FFF2CC")  # derived rows
NP_FILL = PatternFill("solid", fgColor="F2F2F2")   # not-populated / note rows
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")


def style_header(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = BORDER
    ws.row_dimensions[row].height = 30
    # Use a string reference (not ws.cell(...)) so we do not materialize an
    # empty row below the header before data is appended.
    ws.freeze_panes = "A%d" % (row + 1)


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ---------------------------------------------------------------------------
# Source model, extracted from CASPCFNJ
# ---------------------------------------------------------------------------
# Output file 1 : SRCPCF-OUT  (FD SRCPCF-OUT / 01 CLMO-RECORD PIC X(754)),
#                 written FROM WS-SRCPCF-OUT at line 732.
#   WS-SRCPCF-OUT =  PRO-CLMPRFX (copybook CLMPREFX, prefix PRO)   [line 218]
#                 +  CLM-HMS-601 (copybook FDPCF601, prefix CLM)   [line 224]
# Output file 2 : CASE-PCF-MATCH (FD / 01 MATCH-RECORD PIC X(80)),
#                 written FROM WS-MATCH-OUT - layout defined inline (lines 226-235).

# --- PRO- prefix fields (output file 1, CLMPREFX) --------------------------
# (target, source, source_copybook, line, mapping_type, note)
PRO_FIELDS = [
    ("PRO-RECIPIENT-ID-NUM", "CLMI-PCF-MA-NUM", "FDPCF602 (CLMI)", 543,
     "direct MOVE (renamed)", "Recipient Medicaid number copied from claim PCF-MA-NUM."),
    ("PRO-HMS-CASE-KEY", "CASET-HMS-CASE-KEY (SUB-I)", "NCTCASE (CASET)", 545,
     "direct MOVE (from case table)", "HMS case key of the matched open case (table occurrence SUB-I)."),
    ("PRO-CLIENT-ID", "CASET-HMS-CLIENT-ID (SUB-I)", "NCTCASE (CASET)", 547,
     "direct MOVE (from case table)", "HMS client id of the matched open case (chg 0014)."),
    ("PRO-CREATE-SOURCE", "WS-SAVE-CREATE-SOURCE", "CNTL-CARDS / CARD-DATA", 553,
     "direct MOVE (control card)", "2-char create-source read from SYS004 control card tag '1' (line 430)."),
    ("PRO-ICN", "CLMI-ICN", "FDPCF602 (CLMI)", 554,
     "1x1 identity + conditional override", "Base = claim ICN; conditionally rebuilt as ICN+suffix (WS-ICN-GROUP-19, line 718)."),
    ("PRO-FORMER-ICN", "CLMI-FORMER-ICN", "FDPCF602 (CLMI)", 555,
     "direct MOVE (renamed)", "Former ICN copied from claim."),
    ("PRO-XACTION-STATUS", "CLMI-XACTION-STATUS", "FDPCF602 (CLMI)", 557,
     "direct MOVE (renamed)", "Transaction status copied from claim."),
    ("PRO-CLM-FROM-DATE", "PFX-APP-DATE-OF-SERVICE", "TPLPREFX (PFX)", 559,
     "direct MOVE (renamed)", "Claim 'from' date taken from TPL prefix APP-DATE-OF-SERVICE (also drives ICD version)."),
    ("PRO-CLAIM-TRANS-TYPE", "PFX-NET-CLAIM-TRANS-TYPE", "TPLPREFX (PFX)", 561,
     "direct MOVE (renamed)", "Net claim transaction type from TPL prefix."),
    ("PRO-INCIDENT-DATE", "WS-INCIDENT-DATE", "derived (4000-FORMAT-DATE)", 562,
     "derived/transformed (date reformat)", "CASET-INCIDENT-DATE UNSTRINGed to YYYYMM, day forced to 01 (lines 985-991)."),
    ("PRO-CLM-THRU-DATE", "WS-CLM-THRU-DATE", "derived (4000-FORMAT-DATE)", 563,
     "derived/transformed (date reformat)", "CASET-CLAIMS-THRU-DATE UNSTRINGed to YYYYMMDD; '99999999' when blank (lines 993-999)."),
]

# --- CLM- base claim fields (output file 1, FDPCF601) : CLMI-x -> CLM-x -----
# 106 fields, straight prefix-rename MOVEs in paragraph 3020 (lines 572-704).
# (target, source, line)
CLM_BASE = [
    ("CLM-PCF-MA-NUM", "CLMI-PCF-MA-NUM", 572),
    ("CLM-PROV-OF-SVC-NUM", "CLMI-PROV-OF-SVC-NUM", 573),
    ("CLM-PROCEDURE-CODE-7", "CLMI-PROCEDURE-CODE-5", 575),
    ("CLM-PROCEDURE-CODE-MOD1", "CLMI-PROCEDURE-CODE-MOD1", 577),
    ("CLM-PROCEDURE-CODE-MOD2", "CLMI-PROCEDURE-CODE-MOD2", 579),
    ("CLM-SURGERY-DATE", "CLMI-SURGERY-DATE", 580),
    ("CLM-PROC-CODE-DENTAL", "CLMI-PROC-CODE-DENTAL", 581),
    ("CLM-CLM-TYPE", "CLMI-CLM-TYPE", 582),
    ("CLM-ORIG-CLM-TYPE", "CLMI-ORIG-CLM-TYPE", 583),
    ("CLM-CATG-OF-SVC", "CLMI-CATG-OF-SVC", 584),
    ("CLM-TYPE-OF-SVC", "CLMI-TYPE-OF-SVC", 585),
    ("CLM-PLACE-OF-SVC", "CLMI-PLACE-OF-SVC", 586),
    ("CLM-DRG", "CLMI-DRG", 587),
    ("CLM-LEVEL-OF-CARE-IND", "CLMI-LEVEL-OF-CARE-IND", 588),
    ("CLM-CLAIM-FORM-IND", "CLMI-CLAIM-FORM-IND", 589),
    ("CLM-PAT-LAST-NAME", "CLMI-PAT-LAST-NAME", 590),
    ("CLM-PAT-FIRST-NAME", "CLMI-PAT-FIRST-NAME", 591),
    ("CLM-PAT-MID-INIT", "CLMI-PAT-MID-INIT", 592),
    ("CLM-PAT-SEX", "CLMI-PAT-SEX", 593),
    ("CLM-PAT-DOB", "CLMI-PAT-DOB", 594),
    ("CLM-XACTION-TYPE", "CLMI-XACTION-TYPE", 595),
    ("CLM-XACTION-STATUS", "CLMI-XACTION-STATUS", 596),
    ("CLM-DOR", "CLMI-DOR", 597),
    ("CLM-REMIT-ID-NUM", "CLMI-REMIT-ID-NUM", 598),
    ("CLM-ADJ-IND", "CLMI-ADJ-IND", 599),
    ("CLM-ADJ-REASON", "CLMI-ADJ-REASON", 600),
    ("CLM-ICN", "CLMI-ICN", 601),
    ("CLM-PRI-DX", "CLMI-PRI-DX", 602),
    ("CLM-SEC-DX", "CLMI-SEC-DX", 603),
    ("CLM-UNIT-VISITS-DAYS", "CLMI-UNIT-VISITS-DAYS", 604),
    ("CLM-CLAIM-FROM-DOS", "CLMI-CLAIM-FROM-DOS", 605),
    ("CLM-CLAIM-THRU-DOS", "CLMI-CLAIM-THRU-DOS", 606),
    ("CLM-RECORD-THRU-DOS", "CLMI-RECORD-THRU-DOS", 607),
    ("CLM-ADMIT-DATE", "CLMI-ADMIT-DATE", 608),
    ("CLM-DISCH-DATE", "CLMI-DISCH-DATE", 609),
    ("CLM-ELIG-INPAT-DAYS", "CLMI-ELIG-INPAT-DAYS", 610),
    ("CLM-INELIG-INPAT-DAYS", "CLMI-INELIG-INPAT-DAYS", 611),
    ("CLM-NATURE-HSP-ADMISSION", "CLMI-NATURE-HSP-ADMISSION", 613),
    ("CLM-PAT-STATUS-CODE", "CLMI-PAT-STATUS-CODE", 614),
    ("CLM-TRAUMA-IND", "CLMI-TRAUMA-IND", 615),
    ("CLM-TPL-ACCIDENT-IND", "CLMI-TPL-ACCIDENT-IND", 616),
    ("CLM-PROV-CLAIM-NUM", "CLMI-PROV-CLAIM-NUM", 617),
    ("CLM-MA-BILLED-HDR", "CLMI-MA-BILLED-HDR", 618),
    ("CLM-OTHER-INS-PAID-HDR", "CLMI-OTHER-INS-PAID-HDR", 620),
    ("CLM-MC-COINS-DTL", "CLMI-MC-COINS-DTL", 621),
    ("CLM-MC-DEDUCT-DTL", "CLMI-MC-DEDUCT-DTL", 622),
    ("CLM-DRG-ALLOWED", "CLMI-DRG-ALLOWED", 623),
    ("CLM-MA-BILLED-NET-HDR", "CLMI-MA-BILLED-NET-HDR", 624),
    ("CLM-TOT-MA-PAID-HDR", "CLMI-TOT-MA-PAID-HDR", 625),
    ("CLM-MC-ALLOWED-DTL", "CLMI-MC-ALLOWED-DTL", 626),
    ("CLM-OTHER-INS-PAID-DTL", "CLMI-OTHER-INS-PAID-DTL", 628),
    ("CLM-MA-BILLED-NET-DTL", "CLMI-MA-BILLED-NET-DTL", 629),
    ("CLM-MA-PAID-DTL", "CLMI-MA-PAID-DTL", 630),
    ("CLM-PROV-TYPE", "CLMI-PROV-TYPE", 631),
    ("CLM-PAY-TO-PROV-NUM", "CLMI-PAY-TO-PROV-NUM", 632),
    ("CLM-REFER-PROV-NUM", "CLMI-REFER-PROV-NUM", 633),
    ("CLM-PROV-SPECIALTY", "CLMI-PROV-SPECIALTY", 634),
    ("CLM-RX-SUPPLY-DAYS", "CLMI-RX-SUPPLY-DAYS", 635),
    ("CLM-RX-REFILL-IND", "CLMI-RX-REFILL-IND", 636),
    ("CLM-VERSION-NUMBER", "CLMI-VERSION-NUMBER", 637),
    ("CLM-PCF-MC-NUM", "CLMI-PCF-MC-NUM", 638),
    ("CLM-MC-ATTACHMENT-IND", "CLMI-MC-ATTACHMENT-IND", 639),
    ("CLM-MC-FORCE-IND", "CLMI-MC-FORCE-IND", 640),
    ("CLM-MC-BENEFITS-EXHAUSTED", "CLMI-MC-BENEFITS-EXHAUSTED", 642),
    ("CLM-FAMILY-PLANNING-IND", "CLMI-FAMILY-PLANNING-IND", 644),
    ("CLM-MULTI-MATCH-IND", "CLMI-MULTI-MATCH-IND", 645),
    ("CLM-RUNTIME-SEQ-NUM", "CLMI-RUNTIME-SEQ-NUM", 646),
    ("CLM-NUM-REV-CHARGES", "CLMI-NUM-REV-CHARGES", 647),
    ("CLM-NUM-OTH-SEGS", "CLMI-NUM-OTH-SEGS", 648),
    ("CLM-FILE-DATA-SOURCE", "CLMI-FILE-DATA-SOURCE", 649),
    ("CLM-PCF-CONTRACT-NUM", "CLMI-PCF-CONTRACT-NUM", 650),
    ("CLM-CLAIM-EMERGENCY-IND", "CLMI-CLAIM-EMERGENCY-IND", 652),
    ("CLM-STAT-DOLLARS-HDR-OR-DTL", "CLMI-STAT-DOLLARS-HDR-OR-DTL", 654),
    ("CLM-SUMMED-BILL-ANC-DOLLARS", "CLMI-SUMMED-BILL-ANC-DOLLARS", 656),
    ("CLM-BILL-TYPE", "CLMI-BILL-TYPE", 657),
    ("CLM-RECORD-FROM-DOS-DTL", "CLMI-RECORD-FROM-DOS-DTL", 659),
    ("CLM-FORMER-ICN", "CLMI-FORMER-ICN", 660),
    ("CLM-SRC-LINE-ITEM-DTL-CNT", "CLMI-SRC-LINE-ITEM-DTL-CNT", 662),
    ("CLM-NET-DUPE-DATE", "CLMI-NET-DUPE-DATE", 664),
    ("CLM-OTP-COB-IND", "CLMI-OTP-COB-IND", 666),
    ("CLM-PM-USER-AREA", "CLMI-PM-USER-AREA", 667),
    ("CLM-PM-USER-AREA-2", "CLMI-PM-USER-AREA-2", 668),
    ("CLM-XWALKED-PROC-CODE1", "CLMI-XWALKED-PROC-CODE1", 670),
    ("CLM-APPLIED-INCOME", "CLMI-APPLIED-INCOME", 671),
    ("CLM-NUM-LEAVE-DAYS", "CLMI-NUM-LEAVE-DAYS", 672),
    ("CLM-RATE-PAID", "CLMI-RATE-PAID", 673),
    ("CLM-FACILITY-PROV-NUM", "CLMI-FACILITY-PROV-NUM", 674),
    ("CLM-DX-3", "CLMI-DX-3", 675),
    ("CLM-DX-4", "CLMI-DX-4", 676),
    ("CLM-DX-5", "CLMI-DX-5", 677),
    ("CLM-INST-ADDITIONAL-FIELDS", "CLMI-INST-ADDITIONAL-FIELDS", 679),
    ("CLM-NEW-DRG", "CLMI-NEW-DRG", 680),
    ("CLM-PROV-NUM-FROM-PREFIX", "CLMI-PROV-NUM-FROM-PREFIX", 682),
    ("CLM-DATE-REFORMATTED", "CLMI-DATE-REFORMATTED", 683),
    ("CLM-SYS-DELETE-STATUS", "CLMI-SYS-DELETE-STATUS", 684),
    ("CLM-STAMP-DATE-TIME", "CLMI-STAMP-DATE-TIME", 685),
    ("CLM-STAMP-SEQ-NUM", "CLMI-STAMP-SEQ-NUM", 686),
    ("CLM-ORIG-TOT-MA-PAID-HDR", "CLMI-ORIG-TOT-MA-PAID-HDR", 688),
    ("CLM-ORIG-MA-PAID-DTL", "CLMI-ORIG-MA-PAID-DTL", 690),
    ("CLM-PCF-CLASSIFICATION-STAT", "CLMI-PCF-CLASSIFICATION-STAT", 692),
    ("CLM-PCF-CLASSIFICATION-DATE", "CLMI-PCF-CLASSIFICATION-DATE", 694),
    ("CLM-PCF-IN-PROGRESS-IND", "CLMI-PCF-IN-PROGRESS-IND", 696),
    ("CLM-PCF-LOGICAL-DELETE-IND", "CLMI-PCF-LOGICAL-DELETE-IND", 698),
    ("CLM-PCF-HMS-ICN-SUFFIX", "CLMI-PCF-HMS-ICN-SUFFIX", 700),
    ("CLM-PCF-FILE-NAME", "CLMI-PCF-FILE-NAME", 702),
    ("CLM-PCF-VERSION-NUMBER", "CLMI-PCF-VERSION-NUMBER", 704),
]

# Fields whose FINAL value can be overridden by the diagnosis/procedure
# crosswalk (paragraph 3030-REVIEW-CLMS-VERS) and/or other conditional logic.
# Maps CLM target -> dict of per-input-format source expression.
# Base/other format always runs first in 3020 (lines 572-704).
XWALK = {
    "CLM-PROCEDURE-CODE-7": {
        "base": "CLMI-PROCEDURE-CODE-5 (575)",
        "mama6": "CA-A-PROC-CDE (781)",
        "mama5": "CB-A-PROC-CDE (809)",
        "enct5": "CC-E-PROC-CDE if clm-type<>'12' (889)",
        "enct4": "CD-E-PROC-CDE if clm-type<>'12' (966)",
        "mtype": "crosswalk/translated value",
    },
    "CLM-PROCEDURE-CODE-MOD1": {
        "base": "CLMI-PROCEDURE-CODE-MOD1 (577)",
        "mama6": "SPACES (782)", "mama5": "SPACES (810)",
        "enct5": "SPACES (890)", "enct4": "SPACES (967)",
        "mtype": "direct MOVE + override to SPACES",
    },
    "CLM-PROCEDURE-CODE-MOD2": {
        "base": "CLMI-PROCEDURE-CODE-MOD2 (579)",
        "mama6": "SPACES (783)", "mama5": "SPACES (811)",
        "enct5": "SPACES (891)", "enct4": "SPACES (968)",
        "mtype": "direct MOVE + override to SPACES",
    },
    "CLM-PRI-DX": {
        "base": "CLMI-PRI-DX (602)",
        "mama6": "CA-A-DIAG-CDE(1) (784)", "mama5": "CB-A-DIAG-CDE(1) (812)",
        "enct5": "CC-*-DIAG(1) by clm-type (824..881)",
        "enct4": "CD-*-DIAG(1) by clm-type (900..959)",
        "mtype": "crosswalk/translated value",
    },
    "CLM-SEC-DX": {
        "base": "CLMI-SEC-DX (603)",
        "mama6": "CA-A-DIAG-CDE(2) (785)", "mama5": "CB-A-DIAG-CDE(2) (813)",
        "enct5": "CC-*-DIAG(2) by clm-type", "enct4": "CD-*-DIAG(2) by clm-type",
        "mtype": "crosswalk/translated value",
    },
    "CLM-DX-3": {
        "base": "CLMI-DX-3 (675)",
        "mama6": "CA-A-DIAG-CDE(3) (786)", "mama5": "CB-A-DIAG-CDE(3) (814)",
        "enct5": "CC-*-DIAG(3) by clm-type", "enct4": "CD-*-DIAG(3) by clm-type",
        "mtype": "crosswalk/translated value",
    },
    "CLM-DX-4": {
        "base": "CLMI-DX-4 (676)",
        "mama6": "CA-A-DIAG-CDE(4) (787)", "mama5": "CB-A-DIAG-CDE(4) (815)",
        "enct5": "CC-*-DIAG(4) by clm-type", "enct4": "CD-*-DIAG(4) by clm-type",
        "mtype": "crosswalk/translated value",
    },
    "CLM-DX-5": {
        "base": "CLMI-DX-5 (677)",
        "mama6": "CA-A-DIAG-CDE(5) (788)", "mama5": "CB-A-DIAG-CDE(5) (816)",
        "enct5": "CC-*-DIAG(5) by clm-type", "enct4": "CD-*-DIAG(5) by clm-type",
        "mtype": "crosswalk/translated value",
    },
    "CLM-ICN": {
        "base": "CLMI-ICN (601)",
        "mama6": "WS-ICN-GROUP-19 if ICN-suffix cond (720)",
        "mama5": "WS-ICN-GROUP-19 if ICN-suffix cond (720)",
        "enct5": "WS-ICN-GROUP-19 if ICN-suffix cond (720)",
        "enct4": "WS-ICN-GROUP-19 if ICN-suffix cond (720)",
        "mtype": "1x1 identity + conditional override",
    },
}

# Override-only field: CLM-CDE-ICD-VERSION is never set in the 572-704 block;
# it is derived from the from-date at lines 726-730.
CLM_DERIVED = [
    ("CLM-CDE-ICD-VERSION", "PRO-CLM-FROM-DATE", "derived (726-730)", 727,
     "derived/transformed (date->code)",
     "'09' when PRO-CLM-FROM-DATE < 20151001, else '10' (ICD-9 vs ICD-10 indicator)."),
]

# --- Output file 2 : CASE-PCF-MATCH (80-byte report, layout fully in source) -
# WS-MATCH-OUT (lines 226-235). byte offsets computed from the inline PICs.
# (target, pic, byte_from, byte_to, source, line, mtype, sample_in, sample_out, comment)
MATCH_FIELDS = [
    ("FILLER (F)", "X(03)", 1, 3, "literal ' * '", 227,
     "hardcoded constant", "n/a", " * ", "Row delimiter constant."),
    ("MATCH-CASE-ID-OUT", "X(20)", 4, 23, "WS-CASE-ID <= CASET-RECIPIENT-ID-NUM(1)", 1527,
     "direct MOVE (detail) / 'CASE ID' (header)",
     "CASET-RECIPIENT-ID-NUM(1)='12345678901234567890'",
     "12345678901234567890", "Header row shows literal 'CASE ID' (line 228); detail rows show the matched case id."),
    ("FILLER (F)", "X(03)", 24, 26, "literal ' * '", 229,
     "hardcoded constant", "n/a", " * ", "Column delimiter constant."),
    ("MATCH-TOT-PCF-REC-OUT", "X(19)", 27, 45, "WS-CONVERT-MATCH (edited WS-TOT-PCF-REC-MATCH)", 1533,
     "derived/transformed (numeric edit + left-trim)",
     "WS-TOT-PCF-REC-MATCH=3", "3",
     "Header = 'PCF RECORDS MATCHED' (line 230); detail = count of matched PCF records, ZZZZ9-edited then left-trimmed via INSPECT (1529-1534)."),
    ("FILLER (F)", "X(03)", 46, 48, "literal ' * '", 232,
     "hardcoded constant", "n/a", " * ", "Column delimiter constant."),
    ("MATCH-TOT-PCF-MA-PAID-OUT", "X(18)", 49, 66, "WS-CONVERT-PCF-PAID (edited WS-TOT-PCF-MA-PAID)", 1540,
     "derived/transformed (currency edit + left-trim)",
     "WS-TOT-PCF-MA-PAID=1234.56", "$1,234.56",
     "Header = 'PCF TOT $$ MA PAID' (line 233); detail = running $ total, $$$,$$$,$$$,$$9.99-edited then left-trimmed (1536-1541). On COMP-3 overflow a 'TOO MANY MATCHES' message replaces cols 27+ (1543)."),
    ("FILLER (F)", "X(02)", 67, 68, "literal ' *'", 235,
     "hardcoded constant", "n/a", " *", "Trailing delimiter constant (bytes 69-80 remain spaces in the X(80) record)."),
]


# ---------------------------------------------------------------------------
# Inferred PIC map (only where program logic makes type/length determinable).
# Everything else -> unknown (copybook not in repo).
# ---------------------------------------------------------------------------
INFERRED_PIC = {
    "CLM-PRI-DX": "X(7) [inferred: 3+4 expanded diag]",
    "CLM-SEC-DX": "X(7) [inferred]",
    "CLM-DX-3": "X(7) [inferred]",
    "CLM-DX-4": "X(7) [inferred]",
    "CLM-DX-5": "X(7) [inferred]",
    "CLM-PROCEDURE-CODE-7": "X(7) [inferred: 'increased size' proc]",
    "CLM-CDE-ICD-VERSION": "X(2) [inferred: '09'/'10']",
    "CLM-CLAIM-FROM-DOS": "9(7) [inferred: 0YYMMDD, cmp 0040000/0900000 @464-465]",
    "PRO-CLM-FROM-DATE": "9(8) [inferred: cmp 20151001 @726]",
    "PRO-INCIDENT-DATE": "9(8)/X(8) [inferred: YYYYMMDD]",
    "PRO-CLM-THRU-DATE": "9(8)/X(8) [inferred: YYYYMMDD or 99999999]",
}
UNKNOWN_PIC = "[FDPCF601 not in repo]"
UNKNOWN_PIC_PRO = "[CLMPREFX not in repo]"


# ---------------------------------------------------------------------------
# Illustrative sample value generator (marked *illustrative*; copybooks absent)
# ---------------------------------------------------------------------------
def sample_for(name):
    """Return an illustrative value consistent with the field's apparent domain.

    Copybook PICs are unavailable, so these are explicitly illustrative.
    """
    n = name.upper()
    if "DIAG" in n or n.endswith("-DX") or "-DX-" in n or "PRI-DX" in n or "SEC-DX" in n:
        return "E119   "            # ICD code padded to 7
    if "PROCEDURE-CODE-7" in n or "PROC-CDE" in n or "PROC-CODE" in n:
        return "0360T  "            # CPT/HCPCS padded to 7
    if "ICD-VERSION" in n:
        return "09"
    if "MOD1" in n or "MOD2" in n:
        return "  "                 # modifier spaces
    if "DOB" in n:
        return "19600115"
    if "FROM-DOS" in n or "THRU-DOS" in n:
        return "0150315"            # 0YYMMDD
    if "DATE" in n or "-DOS" in n or n.endswith("-DOR") or "DOR" in n:
        return "20150315"
    if any(k in n for k in ("PAID", "BILLED", "DOLLARS", "ALLOWED", "INCOME",
                             "DEDUCT", "COINS", "RATE", "AMT")):
        return "0000123456"         # money (implied decimals)
    if "SEX" in n:
        return "F"
    if "LAST-NAME" in n:
        return "DOE                 "
    if "FIRST-NAME" in n:
        return "JANE           "
    if "MID-INIT" in n:
        return "A"
    if any(k in n for k in ("DAYS", "VISITS", "CNT", "-CTR", "SEQ", "NUM-REV",
                             "NUM-OTH", "NUM-LEAVE", "DRG")):
        return "00005"
    if "ICN" in n:
        return "20151234567890123"  # 17-char ICN
    if "CASE-KEY" in n:
        return "CASE0001234567890   "
    if "CLIENT-ID" in n:
        return "CLIENT0001     "
    if "MA-NUM" in n or "MEDICAID" in n or "RECIPIENT-ID" in n:
        return "12345678901234567890"
    if "MC-NUM" in n:
        return "MC12345678   "
    if "PROV" in n and "NUM" in n:
        return "09999996"
    if "CONTRACT-NUM" in n:
        return "0032600"
    if "USER-AREA" in n:
        return "PB                         "
    if "CLM-TYPE" in n or "TYPE" in n or "CODE" in n or "IND" in n or "STATUS" in n or "FLAG" in n:
        return "01"
    return "SAMPLE         "


def pic_for(name, is_pro=False):
    if name in INFERRED_PIC:
        return INFERRED_PIC[name]
    return UNKNOWN_PIC_PRO if is_pro else UNKNOWN_PIC


# ---------------------------------------------------------------------------
# Sheet 1 : Field Mapping
# ---------------------------------------------------------------------------
def build_sheet_field_mapping(wb):
    ws = wb.active
    ws.title = "Field Mapping"
    headers = [
        "Field No", "Output Copybook", "Target Field", "Picture", "Byte Location",
        "Mapping Type", "Source Copybook", "Source Field", "Sample Input", "Sample Output",
        "Src: MAMA v06 (FDNJTPC5/CA)", "Src: MAMA/CCAR v05 (FDNJTPC4/CB)",
        "Src: ENCT v05 (FDNJ5INS/CC)", "Src: ENCT v04 (FDNJ4INS/CD)",
        "Src: Base/other (FDPCF602)", "Output Path", "Comments",
    ]
    ws.append(headers)
    style_header(ws, len(headers))

    fno = 0
    row_meta = []  # (excel_row, fill)

    def emit(out_cb, target, pic, byteloc, mtype, src_cb, src, s_in, s_out,
             m6, m5, e5, e4, base, path, comment, fill=None):
        nonlocal fno
        fno += 1
        ws.append([fno, out_cb, target, pic, byteloc, mtype, src_cb, src, s_in, s_out,
                   m6, m5, e5, e4, base, path, comment])
        row_meta.append((ws.max_row, fill))

    # ---- Output file 1 : SRCPCF-OUT - PRO- prefix ----
    for (target, src, src_cb, line, mtype, note) in PRO_FIELDS:
        s = sample_for(target)
        is_derived = "derived" in mtype
        fill = DRV_FILL if is_derived else None
        if target == "PRO-ICN":
            m = base = "CLMI-ICN (554); ICN+suffix via WS-ICN-GROUP-19 (718)"
            fill = XW_FILL
        else:
            m = base = f"{src} ({line})"
        # transformed outputs
        if target == "PRO-INCIDENT-DATE":
            s_in, s_out = "CASET-INCIDENT-DATE='2015-03' *illustrative*", "20150301 *illustrative*"
        elif target == "PRO-CLM-THRU-DATE":
            s_in, s_out = "CASET-CLAIMS-THRU-DATE='2016-06-30'", "20160630 *illustrative*"
        elif target == "PRO-CREATE-SOURCE":
            s_in, s_out = "CARD-DATA='NJ'", "NJ"
        else:
            s_in = s_out = f"{s} *illustrative*"
        emit("CLMPREFX (PRO)", target, pic_for(target, is_pro=True),
             "n/a - CLMPREFX not in repo", mtype, src_cb, src, s_in, s_out,
             m, m, m, m, base, "SRCPCF-OUT (claim extract)", note, fill)

    # ---- Output file 1 : SRCPCF-OUT - CLM- claim area (FDPCF601) ----
    for order, (target, src, line) in enumerate(CLM_BASE, start=1):
        s = sample_for(target)
        if target in XWALK:
            xw = XWALK[target]
            mtype = xw["mtype"]
            m6, m5, e5, e4, base = xw["mama6"], xw["mama5"], xw["enct5"], xw["enct4"], xw["base"]
            fill = XW_FILL
            comment = ("Base move CLMI->CLM (line %d); final value overridden by the "
                       "diagnosis/procedure crosswalk in 3030-REVIEW-CLMS-VERS per input "
                       "file type/version/claim type." % line)
            if target == "CLM-ICN":
                comment = ("Base = CLMI-ICN (601); conditionally rebuilt as 17-char ICN + "
                           "2-char suffix (WS-ICN-GROUP-19) at line 720.")
            s_in, s_out = f"{sample_for(src)} *illustrative*", f"{s} *illustrative*"
        else:
            mtype = "1x1 identity"
            m6 = m5 = e5 = e4 = base = f"{src} ({line})"
            fill = None
            comment = "Prefix rename CLMI->CLM; value passes through unchanged."
            s_in = s_out = f"{s} *illustrative*"
        emit("FDPCF601 (CLM)", target, pic_for(target),
             "n/a - FDPCF601 not in repo (rel order %d)" % order, mtype,
             "FDPCF602 (CLMI)", src, s_in, s_out,
             m6, m5, e5, e4, base, "SRCPCF-OUT (claim extract)", comment, fill)

    # ---- Output file 1 : derived CLM-CDE-ICD-VERSION ----
    for (target, src, src_cb, line, mtype, note) in CLM_DERIVED:
        emit("FDPCF601 (CLM)", target, pic_for(target),
             "n/a - FDPCF601 not in repo", mtype, src_cb, src,
             "PRO-CLM-FROM-DATE=20150315", "09 *illustrative*",
             "'09'/'10' by from-date", "'09'/'10' by from-date",
             "'09'/'10' by from-date", "'09'/'10' by from-date",
             "'09' if <20151001 else '10'", "SRCPCF-OUT (claim extract)", note, DRV_FILL)

    # ---- Output file 2 : CASE-PCF-MATCH report ----
    for (target, pic, bfrom, bto, src, line, mtype, s_in, s_out, comment) in MATCH_FIELDS:
        fill = NP_FILL if target.startswith("FILLER") else DRV_FILL
        emit("CASE-PCF-MATCH (inline WS-MATCH-OUT)", target, pic,
             "%d-%d" % (bfrom, bto), mtype, "n/a (defined in program)", src,
             s_in, s_out, "n/a", "n/a", "n/a", "n/a", "n/a",
             "CASE-PCF-MATCH (summary report)", comment, fill)

    # ---- Note row : unpopulated copybook fields ----
    emit("FDPCF601 / CLMPREFX", "(remaining copybook fields)", UNKNOWN_PIC,
         "n/a", "not populated / unknown", "n/a", "n/a", "n/a", "n/a",
         "n/a", "n/a", "n/a", "n/a", "n/a", "SRCPCF-OUT (claim extract)",
         "The FDPCF601/CLMPREFX copybooks define the full 754-byte record and may "
         "contain additional elementary/FILLER fields that the program never MOVEs "
         "to. They cannot be enumerated because the copybooks are not in this repo "
         "(INITIALIZE PRO-CLMPRFX @542 and unset CLM fields default to SPACES/ZEROS).",
         NP_FILL)

    # borders + wrapping + per-row fills
    for r, fill in row_meta:
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.alignment = WRAP
            if fill is not None:
                cell.fill = fill

    set_widths(ws, [8, 20, 30, 28, 26, 30, 18, 28, 26, 26, 26, 26, 26, 26, 24, 24, 50])
    return ws


# ---------------------------------------------------------------------------
# Sheet 2 : Overrides & Crosswalks
# ---------------------------------------------------------------------------
def build_sheet_overrides(wb):
    ws = wb.create_sheet("Overrides & Crosswalks")
    headers = ["Path/Format", "Source Field", "Target Field",
               "Assigned/Translated Value", "Condition", "Sample In -> Out", "Description"]
    ws.append(headers)
    style_header(ws, len(headers))

    rows = [
        ("SYS004 control card", "CARD-DATA(1:2)", "PRO-CREATE-SOURCE",
         "2-char create source", "CARD-TAG = '1' (line 428-430)",
         "'NJ' -> 'NJ'",
         "WS-SAVE-CREATE-SOURCE is loaded from the control card tag '1' and stamped onto every output record."),
        ("All paths", "CLMI-ICN + CLMI-PCF-HMS-ICN-SUFFIX", "PRO-ICN / CLM-ICN",
         "19-char ICN = 17-char ICN + 2-char suffix (WS-ICN-GROUP-19)",
         "PM-USER-AREA(1:2) in ('PB','DT') AND PCF-CONTRACT-NUM='0032600' "
         "AND PM-USER-AREA(27:3)<>'DSS' AND HMS-ICN-SUFFIX numeric AND <>0 (706-723)",
         "ICN='20151234567890123' + '07' -> '2015123456789012307'",
         "ICN reformat that appends the HMS ICN suffix for PB/DT user-area claims on contract 0032600."),
        ("All paths", "PRO-CLM-FROM-DATE", "CLM-CDE-ICD-VERSION",
         "'09' (ICD-9) or '10' (ICD-10)",
         "from-date < 20151001 -> '09' else '10' (726-730)",
         "20150315 -> '09';  20151002 -> '10'",
         "ICD version indicator derived from the claim service date."),
        ("Claim data fix (all)", "CLMI-PAY-TO-PROV-NUM", "CLMI-PROV-OF-SVC-NUM -> CLM-PROV-OF-SVC-NUM",
         "pay-to provider number",
         "PROV-OF-SVC-NUM='09999996' AND PCF-CONTRACT-NUM='0032600' (564-568)",
         "'09999996' -> pay-to-prov",
         "When the service provider is the placeholder 09999996 on contract 0032600, the pay-to provider is substituted before output."),
        ("MAMA v06 (FDNJTPC5 / CA)",
         "CA-A-PROC-CDE; CA-A-DIAG-CDE(1..5)", "CLM-PROCEDURE-CODE-7; CLM-PRI-DX..CLM-DX-5",
         "client-side proc code + up to 5 diagnosis codes",
         "PFX-SYS-HMS-ASSIGN-FILE='MAMA' AND PFX-SYS-VERSION='06' AND clm-type in "
         "(01,02,03,04,06,07,09,13,14,15,18,20,21,23) (764-789)",
         "CA-A-DIAG-CDE(1)='E119' -> CLM-PRI-DX='E119'",
         "Institutional NJ TPC v5 layout overlay; MOD1/MOD2 cleared to spaces."),
        ("MAMA/CCAR v05 (FDNJTPC4 / CB)",
         "CB-A-PROC-CDE; CB-A-DIAG-CDE(1..5)", "CLM-PROCEDURE-CODE-7; CLM-PRI-DX..CLM-DX-5",
         "client-side proc code + up to 5 diagnosis codes",
         "PFX-SYS-HMS-ASSIGN-FILE in ('MAMA','CCAR') AND PFX-SYS-VERSION='05' AND clm-type in list (791-817)",
         "CB-A-DIAG-CDE(1)='I10' -> CLM-PRI-DX='I10'",
         "Claims-activity v4 layout overlay; MOD1/MOD2 cleared to spaces."),
        ("ENCT v05 (FDNJ5INS / CC)",
         "CC-HNI/OP/PR/HH/TR/VI-DIAG(1..5); CC-E-PROC-CDE",
         "CLM-PRI-DX..CLM-DX-5; CLM-PROCEDURE-CODE-7",
         "diagnosis set chosen by claim type; encounter proc code",
         "PFX-SYS-HMS-ASSIGN-FILE='ENCT' AND PFX-SYS-VERSION='05'; EVALUATE CLMI-CLM-TYPE "
         "(01,02,03,04,06,07,18) (819-892)",
         "type 02: CC-HNI-DIAG(1) if not spaces else CC-OP-DIAG(1) -> CLM-PRI-DX",
         "Encounter v5; type 02 falls back from HNI to OP diag; proc set for clm-type<>'12'."),
        ("ENCT v04 (FDNJ4INS / CD)",
         "CD-HNI/OP/PR/HH/TR/VI-DIAG(1..5); CD-E-PROC-CDE",
         "CLM-PRI-DX..CLM-DX-5; CLM-PROCEDURE-CODE-7",
         "diagnosis set chosen by claim type; encounter proc code",
         "PFX-SYS-HMS-ASSIGN-FILE='ENCT' AND PFX-SYS-VERSION='04'; EVALUATE CLMI-CLM-TYPE (895-969)",
         "type 04: CD-PR-DIAG(1) -> CLM-PRI-DX (note: PR type omits DIAG(2)->SEC-DX, line 942-945)",
         "Encounter v4; proc set for clm-type<>'12'."),
        ("Other file/version", "n/a (ADD 1 TO OTHER-VERS)", "n/a",
         "no diag/proc overlay", "none of the above match (5500-PROCESS-RECS, 970-972/1518-1523)",
         "n/a",
         "Record is still written with the base CLMI->CLM moves; only the OTHER-VERS counter increments."),
        ("CASE-PCF-MATCH report", "WS-TOT-PCF-REC-MATCH; WS-TOT-PCF-MA-PAID",
         "MATCH-TOT-PCF-REC-OUT; MATCH-TOT-PCF-MA-PAID-OUT",
         "edited count and currency total",
         "written per matched case id, or 'TOO MANY MATCHES' on COMP-3 size error (1528-1546)",
         "3 -> '3';  1234.56 -> '$1,234.56'",
         "Summary totals accumulated per case id and emitted to the match report."),
    ]
    for row in rows:
        ws.append(list(row))
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.alignment = WRAP
            cell.fill = XW_FILL
    set_widths(ws, [26, 34, 34, 30, 48, 34, 52])
    return ws


# ---------------------------------------------------------------------------
# Sheet 3 : Routing / Format Logic
# ---------------------------------------------------------------------------
def build_sheet_routing(wb):
    ws = wb.create_sheet("Routing - Format Logic")
    ws.append(["CASPCFNJ - Routing / Format Logic"])
    ws["A1"].font = Font(bold=True, size=13, color="1F4E78")
    ws.append([])

    blocks = [
        ("Program role",
         "Reformat + Multi-output + record selection (filter). One PCF claim input (plus a "
         "case master and control cards) is reformatted field-by-field from the HMS-600 input "
         "layout (FDPCF602) into the HMS-601 output layout (FDPCF601, 'increased size diag "
         "code'), prefixed with CLMPREFX, and written to SRCPCF-OUT. A second output "
         "(CASE-PCF-MATCH) is a human-readable match summary report."),
        ("Inputs",
         "SRCPCF-IN  - PCF claim file, VARYING 4..32752, read INTO WS-CLMI-RECORD (1500-READ, line 399).\n"
         "CASEFL-IN  - CAS2000 case master, F 384, read INTO WS-CASE-RECORD, loaded to table "
         "WS-CASE-TABLE OCCURS 40 (1600-READ / 3010-LOAD, lines 408-527).\n"
         "CNTL-CARDS - SYS004 control cards, F 80; tag '1' sets the create-source (1650-READ, 422-430)."),
        ("Outputs (2 paths)",
         "1) SRCPCF-OUT  (CLMO-RECORD PIC X(754)) - the reformatted PCF claim extract; one output "
         "record per matched/in-window claim (WRITE at line 732).\n"
         "2) CASE-PCF-MATCH (MATCH-RECORD PIC X(80)) - header + one detail line per matched case id "
         "+ trailer (WRITEs at 437, 1547, 1711-1716)."),
        ("How many outputs per input",
         "At most ONE SRCPCF-OUT record is produced per qualifying input claim. A claim produces "
         "zero output records if it is skipped/filtered or falls outside the matched case's date "
         "window. The CASE-PCF-MATCH report emits one detail line per matched case id (aggregated "
         "across that case's claims), not per claim."),
        ("Selection: claim must be processed",
         "2000-MAINLINE only continues when PFX-SYS-EXIT-FROM-REF-STATUS = 'Y' (line 452); "
         "otherwise the claim is skipped and the next PCF record is read."),
        ("Match condition",
         "The claim's PFX-APP-MEDICAID-NO must line up with a case recipient id "
         "(CASE-RECIPIENT-ID-NUM / CASET-RECIPIENT-ID-NUM). The case file is advanced until "
         "CASE-RECIPIENT-ID-NUM >= PFX-APP-MEDICAID-NO (445-447); matched recipients are loaded "
         "into the 40-entry case table and read back (3010/3020)."),
        ("Write condition (date window + open case)",
         "In 3020-READ-CASE-TABLE a claim is written only when the case is OPEN "
         "(CASET-CASE-STATUS-CODE = 'O' or X'96', line 537-538) AND "
         "PFX-APP-DATE-OF-SERVICE >= WS-INCIDENT-DATE-NEW-RE (incident date, day=01) AND "
         "<= WS-CLM-THRU-DATE-N (claims-thru date) (540-541)."),
        ("Skip / reject conditions",
         "A) DSS skip: PCF-CONTRACT-NUM='0032600' AND PM-USER-AREA(27:3)='DSS' AND "
         "CLAIM-FROM-DOS in (0040000,0900000) -> ADD REC-SKIP-DSS, read next (457-469).\n"
         "B) Placeholder provider skip: PROV-OF-SVC-NUM='09999996' AND PAY-TO-PROV-NUM='09999996' "
         "-> ADD REC-SKIP-PROV, read next (471-476).\n"
         "C) EXIT-FROM-REF-STATUS <> 'Y' -> skip (452-455).\n"
         "D) Out-of-window / non-open case -> not written (537-541)."),
        ("Format routing (diag/proc crosswalk)",
         "3030-REVIEW-CLMS-VERS routes by PFX-SYS-HMS-ASSIGN-FILE + PFX-SYS-VERSION:\n"
         "  MAMA + 06            -> FDNJTPC5 (CA)   [line 764]\n"
         "  MAMA/CCAR + 05       -> FDNJTPC4 (CB)   [line 791]\n"
         "  ENCT + 05            -> FDNJ5INS (CC)   [line 819]\n"
         "  ENCT + 04            -> FDNJ4INS (CD)   [line 895]\n"
         "  anything else        -> 5500-PROCESS-RECS, ADD 1 TO OTHER-VERS [line 971/1518]\n"
         "Within ENCT, the diagnosis source array is further chosen by CLMI-CLM-TYPE via EVALUATE."),
        ("Stop / termination",
         "0000-MAIN loops 2000-MAINLINE UNTIL PCF-EOF, then 9000-TERMINATION. Processing stops at "
         "end of the PCF file - the case file may not be read to completion (lines 1561-1577). "
         "No ABEND/STOP-with-error path; STOP RUN at line 365. COMP-3 overflow on the match totals "
         "sets SIZE-ERROR and prints 'TOO MANY MATCHES' rather than abending."),
    ]
    r = 3
    for title, body in blocks:
        ws.cell(row=r, column=1, value=title).font = Font(bold=True, color="1F4E78")
        ws.cell(row=r, column=1).fill = SUB_FILL
        ws.cell(row=r, column=1).alignment = TOP
        cell = ws.cell(row=r, column=2, value=body)
        cell.alignment = WRAP
        ws.row_dimensions[r].height = max(30, 14 * (body.count("\n") + 2))
        r += 1
    set_widths(ws, [30, 110])
    return ws


# ---------------------------------------------------------------------------
# Sheet 4 : Sample Records
# ---------------------------------------------------------------------------
def build_sheet_samples(wb):
    ws = wb.create_sheet("Sample Records")
    ws.append(["CASPCFNJ - Worked Sample (one input -> matching outputs)"])
    ws["A1"].font = Font(bold=True, size=13, color="1F4E78")
    ws.append(["All values are *illustrative* - copybook PICs are not in the repo, so lengths/types are inferred."])
    ws["A2"].font = Font(italic=True, color="808080")
    ws.append([])

    def section(title):
        nonlocal_r[0] += 1
        c = ws.cell(row=nonlocal_r[0], column=1, value=title)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HDR_FILL
        ws.merge_cells(start_row=nonlocal_r[0], start_column=1, end_row=nonlocal_r[0], end_column=3)

    def kv(k, v, note=""):
        nonlocal_r[0] += 1
        ws.cell(row=nonlocal_r[0], column=1, value=k).alignment = TOP
        ws.cell(row=nonlocal_r[0], column=2, value=v).alignment = WRAP
        ws.cell(row=nonlocal_r[0], column=3, value=note).alignment = WRAP

    nonlocal_r = [3]

    section("INPUT 1 - SYS004 control card (CNTL-REC X(80))")
    kv("CARD-TAG", "1", "col 1; tag '1' row")
    kv("CARD-DATA (col 35-36)", "NJ", "-> WS-SAVE-CREATE-SOURCE (line 430)")

    section("INPUT 2 - CASEFL-IN case record (CASE-RECORD X(384), NCTCASE) - matched open case")
    kv("CASE-RECIPIENT-ID-NUM", "12345678901234567890", "20-char Medicaid id (matches claim)")
    kv("CASE-CASE-STATUS-CODE", "O", "open case (line 537)")
    kv("CASET-HMS-CASE-KEY", "CASE0001234567890", "-> PRO-HMS-CASE-KEY")
    kv("CASET-HMS-CLIENT-ID", "CLIENT0001", "-> PRO-CLIENT-ID")
    kv("CASET-INCIDENT-DATE", "2015-03", "-> WS-INCIDENT-DATE, day forced 01 -> 20150301")
    kv("CASET-CLAIMS-THRU-DATE", "2016-06-30", "-> WS-CLM-THRU-DATE -> 20160630")

    section("INPUT 3 - SRCPCF-IN PCF claim (WS-CLMI-RECORD: TPLPREFX 127 + FDPCF602 600 + client 32025)")
    kv("PFX-SYS-EXIT-FROM-REF-STATUS", "Y", "required to process (line 452)")
    kv("PFX-APP-MEDICAID-NO", "12345678901234567890", "matches the case recipient id")
    kv("PFX-APP-DATE-OF-SERVICE", "20150315", "in window [20150301 .. 20160630]; <20151001 -> ICD '09'")
    kv("PFX-NET-CLAIM-TRANS-TYPE", "O", "-> PRO-CLAIM-TRANS-TYPE")
    kv("PFX-SYS-HMS-ASSIGN-FILE / VERSION", "MAMA / 06", "routes diag/proc to FDNJTPC5 (CA) - line 764")
    kv("CLMI-PCF-MA-NUM", "12345678901234567890", "-> PRO-RECIPIENT-ID-NUM & CLM-PCF-MA-NUM")
    kv("CLMI-ICN", "20151234567890123", "17-char ICN -> PRO-ICN / CLM-ICN")
    kv("CLMI-PCF-HMS-ICN-SUFFIX", "07", "appended only under PB/DT suffix rule (706-723)")
    kv("CLMI-CLM-TYPE", "01", "institutional; in the MAMA v06 crosswalk list")
    kv("CLMI-PCF-CONTRACT-NUM", "0032600", "")
    kv("CLMI-PM-USER-AREA(1:2)", "PB", "triggers ICN-suffix reformat")
    kv("CLMI-TOT-MA-PAID-HDR", "1234.56", "accumulated into WS-TOT-PCF-MA-PAID")
    kv("CA-A-PROC-CDE", "0360T", "-> CLM-PROCEDURE-CODE-7 (crosswalk, line 781)")
    kv("CA-A-DIAG-CDE(1..2)", "E119 / I10", "-> CLM-PRI-DX / CLM-SEC-DX (784-785)")

    section("OUTPUT 1 - SRCPCF-OUT (CLMO-RECORD X(754)) - key populated fields")
    kv("PRO-RECIPIENT-ID-NUM", "12345678901234567890", "= CLMI-PCF-MA-NUM (543)")
    kv("PRO-HMS-CASE-KEY", "CASE0001234567890", "= CASET-HMS-CASE-KEY (545)")
    kv("PRO-CLIENT-ID", "CLIENT0001", "= CASET-HMS-CLIENT-ID (547)")
    kv("PRO-CREATE-SOURCE", "NJ", "= WS-SAVE-CREATE-SOURCE (553)")
    kv("PRO-ICN", "2015123456789012307", "ICN 17 + suffix 07 (718)")
    kv("PRO-CLM-FROM-DATE", "20150315", "= PFX-APP-DATE-OF-SERVICE (559)")
    kv("PRO-INCIDENT-DATE", "20150301", "reformatted, day=01 (562/990-991)")
    kv("PRO-CLM-THRU-DATE", "20160630", "reformatted (563/993-998)")
    kv("CLM-PCF-MA-NUM", "12345678901234567890", "1x1 identity (572)")
    kv("CLM-ICN", "2015123456789012307", "ICN+suffix override (720)")
    kv("CLM-CLM-TYPE", "01", "1x1 identity (582)")
    kv("CLM-PROCEDURE-CODE-7", "0360T", "crosswalk from CA-A-PROC-CDE (781)")
    kv("CLM-PROCEDURE-CODE-MOD1/2", "(spaces)", "cleared in crosswalk (782-783)")
    kv("CLM-PRI-DX", "E119", "crosswalk from CA-A-DIAG-CDE(1) (784)")
    kv("CLM-SEC-DX", "I10", "crosswalk from CA-A-DIAG-CDE(2) (785)")
    kv("CLM-TOT-MA-PAID-HDR", "1234.56", "1x1 identity (625)")
    kv("CLM-CDE-ICD-VERSION", "09", "derived: from-date 20150315 < 20151001 (726-727)")
    kv("...~100 other CLM-* fields", "= their CLMI-* source", "1x1 identity pass-through (573-704)")

    section("OUTPUT 2 - CASE-PCF-MATCH (MATCH-RECORD X(80)) - detail line for this case")
    kv("bytes 1-3  FILLER", " * ", "constant")
    kv("bytes 4-23 MATCH-CASE-ID-OUT", "12345678901234567890", "= WS-CASE-ID (1527)")
    kv("bytes 24-26 FILLER", " * ", "constant")
    kv("bytes 27-45 MATCH-TOT-PCF-REC-OUT", "1", "count of matched PCF recs (ZZZZ9 trimmed)")
    kv("bytes 46-48 FILLER", " * ", "constant")
    kv("bytes 49-66 MATCH-TOT-PCF-MA-PAID-OUT", "$1,234.56", "running $ total (trimmed)")
    kv("bytes 67-68 FILLER", " *", "constant")

    for row in range(3, ws.max_row + 1):
        for col in range(1, 4):
            ws.cell(row=row, column=col).border = BORDER
    set_widths(ws, [34, 40, 58])
    return ws


def main():
    wb = Workbook()
    build_sheet_field_mapping(wb)
    build_sheet_overrides(wb)
    build_sheet_routing(wb)
    build_sheet_samples(wb)
    wb.save(OUT_XLSX)
    print("Wrote %s with sheets: %s" % (OUT_XLSX, ", ".join(wb.sheetnames)))


if __name__ == "__main__":
    main()
